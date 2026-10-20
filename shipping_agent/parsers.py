from pathlib import Path
from email import message_from_bytes
from email.policy import default
from html.parser import HTMLParser
import csv
import hashlib
import re
import subprocess
import pdfplumber
from docx import Document
from openpyxl import load_workbook
from .models import SourceUnit
from .config import Settings, ROOT
from .ocr import remote_ocr, local_image_text

SUPPORTED = {'.eml', '.pdf', '.png', '.jpg', '.jpeg', '.tif', '.tiff', '.xlsx', '.xls', '.docx', '.doc', '.csv', '.txt'}


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__(); self.parts = []; self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('style', 'script'): self.hidden += 1
        if tag in ('br', 'div', 'p', 'tr'): self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('style', 'script'): self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden: self.parts.append(data)


def html_text(text):
    parser = TextHTML(); parser.feed(text)
    return ''.join(parser.parts)


def parse_file(path: Path, settings: Settings, work: Path, source: str | None = None, depth=0) -> list[SourceUnit]:
    source = source or path.name
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    base = dict(source=source, sha256=digest, kind=path.suffix.lower().lstrip('.'))
    def unit(**kw): return SourceUnit(**base, **kw)
    try:
        if depth > settings.max_depth: raise ValueError('附件递归超过限制')
        if path.stat().st_size > settings.max_file_bytes: raise ValueError('文件超过25MB限制')
        ext = path.suffix.lower()
        if ext not in SUPPORTED: raise ValueError(f'不支持的格式：{ext}')
        if ext == '.eml':
            msg = message_from_bytes(path.read_bytes(), policy=default)
            body = msg.get_body(preferencelist=('plain', 'html'))
            text = body.get_content() if body else ''
            if body and body.get_content_type() == 'text/html': text = html_text(text)
            headers = {key: str(msg.get(key, '')) for key in ('Subject', 'From', 'To', 'Cc', 'Date', 'Message-ID')}
            units = [unit(text='\n'.join(f'{k}: {v}' for k, v in headers.items()) + '\n\n' + text.replace('\r', ''), metadata={'email_headers': headers})]
            html_parts = '\n'.join(p.get_content() for p in msg.walk() if p.get_content_type() == 'text/html' and not p.is_multipart())
            for n, part in enumerate(msg.walk()):
                if part.is_multipart() or part is body: continue
                name = part.get_filename()
                if not name and part.get_content_disposition() != 'attachment': continue
                name = re.split(r'[/\\]', name or f'attachment-{n}')[ -1]
                data = part.get_payload(decode=True) or b''
                child_source = f'{source}!{name}'
                cid = (part.get('Content-ID') or '').strip('<>')
                if cid and f'cid:{cid}' in html_parts and part.get_content_disposition() != 'attachment':
                    units.append(SourceUnit(source=child_source, sha256=hashlib.sha256(data).hexdigest(), kind='inline_image', metadata={'skipped': '邮件内嵌签名图，不按账单解析'}))
                    continue
                folder = work / 'attachments' / digest / str(n); folder.mkdir(parents=True, exist_ok=True)
                child = folder / name; child.write_bytes(data)
                units.extend(parse_file(child, settings, work, child_source, depth + 1))
                if len(units) > settings.max_units: raise ValueError('附件数量超过限制')
            return units
        if ext in {'.pdf', '.png', '.jpg', '.jpeg', '.tif', '.tiff'}:
            if settings.ocr_mode == 'remote':
                text, meta = remote_ocr(path, source, settings, work / 'ocr_cache')
                return [unit(text=text, metadata=meta)]
            if settings.ocr_mode != 'local': raise ValueError('OCR_MODE必须为local或remote')
            if ext != '.pdf':
                text, meta = local_image_text(path)
                if not text.strip(): raise ValueError('图片OCR未识别到文本，需要人工复核')
                return [unit(text=text, metadata=meta)]
            units = []
            with pdfplumber.open(path) as pdf:
                if len(pdf.pages) > settings.max_pages: raise ValueError('PDF超过100页限制')
                for n, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ''
                    meta = {'engine': 'pdfplumber-text-local'}
                    if len(text.strip()) < 20:
                        import pypdfium2 as pdfium
                        with pdfium.PdfDocument(path) as doc:
                            image = doc[n - 1].render(scale=2).to_pil()
                            img_path = work / f'{digest}-{n}.png'; work.mkdir(parents=True, exist_ok=True)
                            image.save(img_path)
                        text, meta = local_image_text(img_path)
                    if not text.strip(): raise ValueError(f'PDF第{n}页未识别到文本')
                    units.append(unit(text=text, page=n, metadata=meta))
            return units
        if ext in {'.xlsx', '.xls', '.csv'}:
            sheets = []
            if ext == '.xlsx':
                wb = load_workbook(path, read_only=True, data_only=True)
                for sheet in wb:
                    if sheet.max_row > 10000 or sheet.max_column > 200: raise ValueError('表格超过10000行/200列限制')
                    sheets.append((sheet.title, [list(row) for row in sheet.values]))
                wb.close()
            elif ext == '.xls':
                import xlrd
                wb = xlrd.open_workbook(path)
                for sheet in wb.sheets():
                    if sheet.nrows > 10000 or sheet.ncols > 200: raise ValueError('表格超过10000行/200列限制')
                    sheets.append((sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)]))
            else:
                sheets = [(path.stem, list(csv.reader(path.read_text(encoding='utf-8-sig').splitlines())))]
            return [unit(text='\n'.join(' | '.join('' if v is None else str(v) for v in row) for row in rows), tables=[[[None if v is None else str(v) for v in row] for row in rows]], metadata={'sheet': title, 'engine': ext[1:]}) for title, rows in sheets]
        if ext == '.docx':
            doc = Document(path)
            tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
            text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(' | '.join(row) for t in tables for row in t)
            return [unit(text=text, tables=tables, metadata={'engine': 'python-docx'})]
        if ext == '.doc':
            # Legacy binary DOC requires a real converter. Never pretend it is DOCX.
            result = subprocess.run(['powershell.exe', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(ROOT / 'scripts' / 'read_doc.ps1'), '-DocumentPath', str(path.resolve())], capture_output=True, encoding='utf-8', timeout=90)
            if result.returncode: raise ValueError('DOC读取失败：需要Windows Microsoft Word；也可先另存为DOCX。' + result.stderr[-500:])
            return [unit(text=result.stdout, metadata={'engine': 'Word-COM-readonly'})]
        return [unit(text=path.read_text(encoding='utf-8-sig'), metadata={'engine': 'text'})]
    except Exception as exc:
        return [unit(issues=[f'{type(exc).__name__}: {exc}'])]
